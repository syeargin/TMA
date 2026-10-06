# Builds the Team setup and Club setup upload templates. Run from the repo root: python3 docs/templates/build-templates.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

NAVY="15294D"; RED="C8323E"
F=lambda **k: Font(name="Arial", **k)
HDR_FILL=PatternFill("solid", fgColor=NAVY); REQ_FILL=PatternFill("solid", fgColor=RED)
EX_FILL=PatternFill("solid", fgColor="EDF0F5"); INPUT_FILL=PatternFill("solid", fgColor="FFFFFF")
thin=Side(style="thin", color="D5DBE5"); BORDER=Border(bottom=thin)
MAXROW=1000

CHOICES={
 "YesNo":["Y","N"],
 "Days":["Sunday","Monday","Tuesday","Wednesday","Thursday","Friday","Saturday"],
 "EventTypes":["Tournament","Event","Deadline"],
 "Repeats":["None","Weekly","Every 2 weeks","Every 3 weeks"],
 "Shirts":["Youth S","Youth M","Youth L","Youth XL","Adult XS","Adult S","Adult M","Adult L","Adult XL","Adult 2XL"],
 "RefGroups":["A","B"],
 "Lists":["Packing checklist","Uniform items"],
 "Programs":["National","Regional","Boys"],
 "DefaultKinds":["Handbook section","Packing checklist","Uniform items"],
}

def add_choices(wb):
    ws=wb.create_sheet("Choices")
    for c,(name,vals) in enumerate(CHOICES.items(),1):
        ws.cell(1,c,name).font=F(bold=True)
        for r,v in enumerate(vals,2): ws.cell(r,c,v).font=F()
        col=get_column_letter(c)
        wb.defined_names[name]=DefinedName(name, attr_text=f"Choices!${col}$2:${col}${len(vals)+1}")
    ws.sheet_state="hidden"

def dv_list(ws, col, source, prompt=None):
    dv=DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                      errorTitle="Pick from the list", error="Choose one of the options in the dropdown.")
    if prompt: dv.promptTitle="Tip"; dv.prompt=prompt; dv.showInputMessage=True
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{MAXROW}")

def dv_custom(ws, col, formula, err):
    dv=DataValidation(type="custom", formula1=formula, allow_blank=True, showErrorMessage=True, errorTitle="Check this value", error=err)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{MAXROW}")

def dv_date(ws, col):
    dv=DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2040,12,31)", allow_blank=True,
                      showErrorMessage=True, errorTitle="Enter a date", error="Enter a date, for example 2026-12-12.")
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{MAXROW}")

def sheet(wb, title, cols, example, rule_kinds):
    """cols: list of (header, required, width, note); rule_kinds: dict header -> ('list', name) | ('email',) | ('date',) | ('money',) | ('hex',) | ('int',)"""
    ws=wb.create_sheet(title)
    for i,(h,req,w,note) in enumerate(cols,1):
        c=ws.cell(1,i,h+(" *" if req else ""))
        c.font=F(bold=True,color="FFFFFF"); c.fill=REQ_FILL if req else HDR_FILL
        c.alignment=Alignment(wrap_text=True, vertical="center")
        if note: c.comment=Comment(note, "Team Hub")
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=32
    for i,v in enumerate(example,1):
        c=ws.cell(2,i,v); c.font=F(italic=True,color="6C7890"); c.fill=EX_FILL
    ws.freeze_panes="A2"
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_title_rows="1:1"
    for i,(h,*_) in enumerate(cols,1):
        col=get_column_letter(i); rk=rule_kinds.get(h)
        for r in range(2, MAXROW+1):
            pass
        if not rk: 
            fmt=None
        elif rk[0]=="list": dv_list(ws,col,f"={rk[1]}", rk[2] if len(rk)>2 else None)
        elif rk[0]=="email": dv_custom(ws,col,f'=AND(ISNUMBER(SEARCH("@",{col}2)),ISNUMBER(SEARCH(".",{col}2)))',"Enter an email address like name@example.com.")
        elif rk[0]=="date": dv_date(ws,col)
        elif rk[0]=="hex": dv_custom(ws,col,f'=AND(LEN({col}2)=7,LEFT({col}2,1)="#")',"Enter a color like #15294D.")
        elif rk[0]=="money":
            dv=DataValidation(type="decimal",operator="between",formula1="0",formula2="100000",allow_blank=True,showErrorMessage=True,error="Enter an amount in dollars, like 400.")
            ws.add_data_validation(dv); dv.add(f"{col}2:{col}{MAXROW}")
        elif rk[0]=="int":
            dv=DataValidation(type="whole",operator="between",formula1="0",formula2="999",allow_blank=True,showErrorMessage=True,error="Enter a whole number.")
            ws.add_data_validation(dv); dv.add(f"{col}2:{col}{MAXROW}")
        elif rk[0]=="teamid":
            dv=DataValidation(type="list",formula1="=Teams!$A$2:$A$1000",allow_blank=True,showErrorMessage=True,errorTitle="Unknown team",error="Use a Team ID from the Teams tab.")
            ws.add_data_validation(dv); dv.add(f"{col}2:{col}{MAXROW}")
        if rk and rk[0] in ("date",):
            for r in range(2,MAXROW+1): ws.cell(r,i).number_format="yyyy-mm-dd"
        if rk and rk[0]=="money":
            for r in range(2,MAXROW+1): ws.cell(r,i).number_format="$#,##0.00"
    for r in range(3, 60):
        for i in range(1,len(cols)+1): ws.cell(r,i).font=F()
    return ws

import datetime as dt
D=lambda s: dt.date.fromisoformat(s)

# ---------- shared column sets ----------
STAFF=[("First name",True,14,None),("Last name",True,14,None),("Email",True,28,"They get access when they sign in or sign up with this address."),
 ("Mobile",False,14,"Shown to families only if 'Show phone to families' is Y."),
 ("Team admin",False,11,"Y = can invite people and set roles. Each team needs at least one."),("Coach",False,9,None),("Coordinator",False,12,None),
 ("Food",False,8,"Plans meals and assigns families."),("Finance",False,9,"Confirms payments and keeps the ledger."),
 ("Show phone to families",False,13,None)]
STAFF_RULES={"Email":("email",),"Team admin":("list","YesNo"),"Coach":("list","YesNo"),"Coordinator":("list","YesNo"),"Food":("list","YesNo"),"Finance":("list","YesNo"),"Show phone to families":("list","YesNo")}
STAFF_EX=["Morgan","Lee","morgan.lee@example.com","555-0101","Y","Y","N","N","N","Y"]

ROSTER=[("Jersey",False,8,"Used to name the ball cart and volleyballs families on the Schedule tab."),("Player first name",True,14,None),("Player last name",True,14,None),
 ("Shirt size",False,11,None),("Town",False,16,None),("Allergies",False,16,"Optional. Only coaches and the food role will see it."),("Ref group",False,9,"A or B: splits the team for ref jobs at tournaments.")]
for n in (1,2,3,4):
    req= n==1
    ROSTER+= [(f"Parent {n} name",req,18,None),(f"Parent {n} email",req,26,"They get access when they sign in or sign up with this address. A parent with two players uses the same email on both rows." if n==1 else None),
              (f"Parent {n} mobile",False,13,None),(f"Parent {n} invite",False,9,"Y (or blank) invites them as this player's parent. N lists them on the roster only.")]
ROSTER_RULES={"Shirt size":("list","Shirts"),"Ref group":("list","RefGroups"),**{f"Parent {n} email":("email",) for n in (1,2,3,4)},**{f"Parent {n} invite":("list","YesNo") for n in (1,2,3,4)}}
ROSTER_EX=["7","Ava","Carter","Adult S","Duluth, GA","None","A","Jen Carter","jen.carter@example.com","555-0102","Y","Ryan Carter","ryan.carter@example.com","555-0103","Y","","","","","","","",""]

PRACTICE=[("Name",True,18,None),("Day",True,12,None),("Start time",True,11,"Like 6:30 PM. Leave the end blank if it isn't set yet."),("End time",False,11,None),
 ("First date",True,12,None),("Last date",False,12,"Blank = end of the season (June 30)."),("Location",False,22,None),("Note",False,30,None)]
PRACTICE_RULES={"Day":("list","Days"),"First date":("date",),"Last date":("date",)}
PRACTICE_EX=["Saturday practice","Saturday","7:30 AM","10:00 AM",D("2026-09-19"),D("2027-06-19"),"A5 Sportsplex","Bring both jerseys"]

SCHED=[("Type",True,12,None),("Title",True,28,None),("Start date",True,12,None),("End date",False,12,"For tournaments over more than one day."),("Time",False,11,None),
 ("Venue",False,24,None),("City",False,16,None),("Division",False,12,None),("Travel",False,8,"Y for tournaments that need hotels or flights."),
 ("Website",False,24,None),("Admission link",False,24,None),("Hotel",False,20,None),("Hotel link",False,24,None),("Hotel block code",False,14,None),("Book hotel by",False,12,None)]
SCHED_FAMILY=[("Ball cart family",False,12,"Jersey number of the player whose family brings the ball cart, or 'Not needed'."),("Volleyballs family",False,12,"Jersey number, or 'Not needed'.")]
SCHED_REPEAT=[("Repeats",False,13,"For team events only; tournaments don't repeat."),("Repeat days",False,14,"Days of the week, like 'Tue, Thu'."),("Repeat until",False,12,None),("Notes",False,36,None)]
SCHED_RULES={"Type":("list","EventTypes"),"Start date":("date",),"End date":("date",),"Travel":("list","YesNo"),"Book hotel by":("date",),"Repeats":("list","Repeats"),"Repeat until":("date",)}
SCHED_EX=["Tournament","Winter Invitational Classic",D("2026-12-12"),D("2026-12-13"),"","Georgia World Congress Center","Atlanta, GA","13 Open","N",
 "https://www.a5tournaments.com","https://sportwrench.com","","","",None]

def readme(wb, title, intro, tabs, extra):
    ws=wb.active; ws.title="Read me"
    ws.column_dimensions["A"].width=24; ws.column_dimensions["B"].width=70; ws.column_dimensions["C"].width=14
    r=1
    ws.cell(r,1,title).font=F(bold=True,size=16,color=NAVY); r+=1
    ws.cell(r,1,intro).font=F(color="47546C"); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
    ws.cell(r,1).alignment=Alignment(wrap_text=True); ws.row_dimensions[r].height=48; r+=2
    ws.cell(r,1,"How to fill it in").font=F(bold=True,size=12,color=NAVY); r+=1
    legend=[("Red header *","Required column. Rows missing it are listed as problems in the upload preview."),
            ("Navy header","Optional column."),
            ("Grey italic row 2","An example showing the format. Type over it or delete it; it is skipped if left unchanged."),
            ("Dropdown cells","Pick from the list. Hover over a header for a note on that column."),
            ("Dates","Any Excel date, or text like 2026-12-12."),
            ("Times","Text like 6:30 PM."),
            ("Leave out","Availability, travel plans, uniform sizes, meal claims and payments: families enter these themselves. Never add passwords.")]+extra
    for k,v in legend:
        ws.cell(r,1,k).font=F(bold=True); c=ws.cell(r,2,v); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical="top")
        ws.cell(r,1).alignment=Alignment(vertical="top"); r+=1
    r+=1
    ws.cell(r,1,"Tabs").font=F(bold=True,size=12,color=NAVY); ws.cell(r,3,"Rows filled").font=F(bold=True,size=12,color=NAVY); r+=1
    for tab,desc,keycol in tabs:
        ws.cell(r,1,tab).font=F(bold=True); c=ws.cell(r,2,desc); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical="top")
        ws.cell(r,1).alignment=Alignment(vertical="top")
        q=f"'{tab}'" if " " in tab else tab
        ws.cell(r,3,f"=COUNTA({q}!{keycol}3:{keycol}{MAXROW})").font=F()
        r+=1
    r+=1
    ws.cell(r,1,"When you upload").font=F(bold=True,size=12,color=NAVY); r+=1
    for t in ["You see a preview first: what will be created, and each problem by tab and row. Nothing is saved until you confirm.",
              "Uploading again updates what's there instead of adding duplicates. Teams match on Team ID, players on jersey or name, events on type, date and title, people on email. Blank cells keep what's saved.",
              "Nothing is removed: players, events and people already on a team but not in the workbook stay, and so do answers, payments and travel plans families entered.",
              "Each person marked for an invite gets the right role, linked to their player, as soon as they sign in or create an account with that email. Team Hub doesn't email them yet, so let them know."]:
        c=ws.cell(r,1,"•"); c.font=F(); c.alignment=Alignment(horizontal="right",vertical="top")
        c=ws.cell(r,2,t); c.font=F(); c.alignment=Alignment(wrap_text=True); r+=1
    ws.sheet_view.showGridLines=False
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True

# ================= TEAM WORKBOOK =================
wb=Workbook(); add_choices(wb)
readme(wb,"Team setup","Fill in this workbook to set up one team in Team Hub: who's on it, its staff and practices, and its season schedule. Club admins upload it from the team's Members page, or from the club page to create a new team.",
 [("Team","One row: the team's name, season and fund settings.","A"),("Staff","One row per coach, coordinator or other staff member.","C"),
  ("Roster","One row per player, with up to four parents.","B"),("Practices","One row per weekly practice.","A"),
  ("Schedule","One row per tournament, team event or deadline.","B"),("Lists","Optional: packing checklist and uniform items.","B")],[])
sheet(wb,"Team",[("Team name",True,18,None),("Season",True,10,"Like 2026-27."),("Age group",True,10,"Like 13U."),("Level / division",False,14,None),
 ("Team code",False,16,"The tournament registration code, shown on game-day pages."),("Dues amount",False,12,"Each family's team fund deposit, in dollars."),
 ("Dues due date",False,13,None),("Dues label",False,24,None),("Meal cost per person",False,12,"For the tournament meal budget."),("Meals per day",False,10,None),("People fed",False,10,"Players plus coaches.")],
 ["A5 13 Tom","2026-27","13U","13 Open","TEAMCODE123",400,D("2026-11-01"),"Initial team fund deposit",20,2,15],
 {"Dues amount":("money",),"Dues due date":("date",),"Meal cost per person":("money",),"Meals per day":("int",),"People fed":("int",)})
sheet(wb,"Staff",STAFF,STAFF_EX,STAFF_RULES)
sheet(wb,"Roster",ROSTER,ROSTER_EX,ROSTER_RULES)
sheet(wb,"Practices",PRACTICE,PRACTICE_EX,PRACTICE_RULES)
sheet(wb,"Schedule",SCHED+SCHED_FAMILY+SCHED_REPEAT,SCHED_EX+["Not needed","Not needed","None","","","Schedules post the Wednesday before."],SCHED_RULES)
sheet(wb,"Lists",[("List",True,20,None),("Item",True,40,None)],["Packing checklist","Knee and elbow pads"],{"List":("list","Lists")})
wb.move_sheet("Choices", offset=len(wb.sheetnames))
wb.save("web/public/templates/Team-Setup-Template.xlsx")

# ================= CLUB WORKBOOK =================
wb=Workbook(); add_choices(wb)
readme(wb,"Club setup","Fill in this workbook to set up a whole club in Team Hub at once: the club's look and links, every team, its staff and roster, practices by age group and the tournaments teams share. Club admins and site owners upload it from the club page.",
 [("Club","One row: the club's name, colors and notes for every team.","A"),("Club links","Links every team sees on Team Info.","A"),
  ("Club admins","People who run the club; they get club-admin invites.","C"),("Teams","One row per team. Every other tab refers to teams by Team ID.","A"),
  ("Staff","One row per person per team. A coach on two teams has two rows.","A"),("Rosters","One row per player, all teams together.","A"),
  ("Practice patterns","One row per weekly practice, for a team or a whole group.","A"),("Shared schedule","Tournaments and events, each listed once for all the teams going.","A"),
  ("Defaults","Optional: the starting handbook and lists for every team.","A")],
 [("Team ID","Lowercase letters, numbers and dashes, like a5-13-tom. It becomes part of the team's web address and can't be changed later."),
  ("Groups","On Practice patterns and Shared schedule you can name a group instead of listing teams: All, All National, All Regional, All Boys, an age group like 13U, or both like 13U National. Or list Team IDs separated by commas.")])
sheet(wb,"Club",[("Club name",True,20,None),("Short name",False,11,"Used in sentences like 'A5 reimburses the fee'."),("Club ID",True,12,"Lowercase, like a5. Can't be changed later."),
 ("Main color",False,11,"Header, buttons and links, like #15294D."),("Accent color",False,11,"The stripe under the header, like #C8323E."),("Notes for every team",False,50,"Shown on every team's Team Info page.")],
 ["A5 Volleyball","A5","a5","#15294D","#C8323E","Club fees are due in August."],{"Main color":("hex",),"Accent color":("hex",)})
sheet(wb,"Club links",[("Label",True,30,None),("Web address",True,50,"Must start with https://")],["Registration and player-parent contract","https://a5volleyball.sprocketsports.com/"],{})
sheet(wb,"Club admins",[("First name",True,14,None),("Last name",True,14,None),("Email",True,30,None)],["Pat","Rivera","pat.rivera@example.com"],{"Email":("email",)})
sheet(wb,"Teams",[("Team ID",True,16,"Lowercase letters, numbers and dashes. Can't be changed later."),("Team name",True,20,None),("Program",True,11,None),("Age group",True,10,"Like 13U."),
 ("Level / division",False,14,None),("Season",True,10,None),("Team code",False,16,None),("Dues amount",False,12,None),("Dues due date",False,13,None),
 ("Copy setup from",False,16,"Optional Team ID to copy the handbook, lists and practices from.")],
 ["a5-13-tom","A5 13-2 Tom","National","13U","13 Open","2026-27","TEAMCODE123",400,D("2026-11-01"),""],
 {"Program":("list","Programs"),"Dues amount":("money",),"Dues due date":("date",)})
tid=("Team ID",True,14,"Must match a Team ID on the Teams tab.")
sheet(wb,"Staff",[tid]+STAFF,["a5-13-tom"]+STAFF_EX,{"Team ID":("teamid",),**STAFF_RULES})
sheet(wb,"Rosters",[tid]+ROSTER,["a5-13-tom"]+ROSTER_EX,{"Team ID":("teamid",),**ROSTER_RULES})
sheet(wb,"Practice patterns",[("Applies to",True,18,"A Team ID, Team IDs separated by commas, or a group like 13U National.")]+PRACTICE,["13U National"]+PRACTICE_EX,PRACTICE_RULES)
sheet(wb,"Shared schedule",[("Teams",True,18,"Team IDs separated by commas, or a group like All 13U.")]+SCHED+SCHED_REPEAT,
 ["All National"]+SCHED_EX+["None","","","Schedules post the Wednesday before."],SCHED_RULES)
sheet(wb,"Defaults",[("Kind",True,18,None),("Title",False,24,"For handbook sections."),("Text or item",True,60,None)],
 ["Handbook section","Attendance","Practices are mandatory. Tell your coach by text if you'll miss one."],{"Kind":("list","DefaultKinds")})
wb.move_sheet("Choices", offset=len(wb.sheetnames))
wb.save("web/public/templates/Club-Setup-Template.xlsx")
print("ok")
