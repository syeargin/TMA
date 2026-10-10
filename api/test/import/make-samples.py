# Fills copies of the upload templates with made-up people for the import tests.
# Run from the repo root: python3 api/test/import/make-samples.py, then node api/test/import/to-json.mjs
# Row 2 (the template's example) is left as it is, so the tests also check that it's skipped.
import datetime as dt
from openpyxl import load_workbook

D = dt.date.fromisoformat
OUT = "api/test/import/"


def fill(ws, rows, start=3):
    for r, vals in enumerate(rows, start):
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)


def roster(jersey, first, last, p1, e1, p2="", e2="", inv2="", shirt="Adult S", ref="A"):
    return [jersey, first, last, shirt, "Roswell, GA", "", ref, p1, e1, "555-0110", "", p2, e2, "", inv2]


# ---------- team ----------
wb = load_workbook("web/public/templates/Team-Setup-Template.xlsx")
fill(wb["Team"], [["Test 15", "2026-27", "15U", "15 Open", "CODE15", 350, D("2026-11-15"), "Deposit", 18, 2, 14]], start=2)
fill(wb["Staff"], [
    ["Dana", "Fox", "dana.fox@example.com", "555-0201", "Y", "Y", "N", "N", "N", "Y"],
    ["Lee", "Park", "lee.park@example.com", "555-0202", "N", "N", "Y", "N", "N", "N"],
    ["Kim", "Ruiz", "kim.ruiz@example.com", "", "N", "N", "N", "N", "Y", "N"],
])
fill(wb["Roster"], [
    roster("3", "Ivy", "Ruiz", "Kim Ruiz", "kim.ruiz@example.com", "Sam Ruiz", "sam.ruiz@example.com", "N"),
    roster("8", "June", "Hale", "Ana Hale", "ana.hale@example.com"),
    roster("12", "Rae", "Cole", "Bo Cole", "bo.cole@example.com", ref="B"),
    roster("15", "Tess", "Cole", "Bo Cole", "bo.cole@example.com", ref="B"),
])
fill(wb["Practices"], [
    ["Tuesday practice", "Tuesday", dt.time(18, 30), dt.time(20, 30), D("2026-10-06"), None, "Every 2 weeks", "Test Gym", "3", "Navy", ""],
    ["Saturday practice", "Sat", "9:00 AM", "11:00 AM", D("2026-10-10"), D("2027-05-29"), "", "Test Gym", "", "Gray", "Both jerseys"],
])
fill(wb["Schedule"], [
    ["Tournament", "Peach Classic", D("2026-12-05"), D("2026-12-06"), "8:00 AM", "Test Center", "Atlanta, GA", "15 Open", "N",
     "https://example.com/peach", "", "", "", "", None, "8", "Not needed", "None", "", None, "Arrive 45 minutes early."],
    ["Event", "Team dinner", D("2026-10-14"), None, "6:00 PM", "Pizza place", "", "", "", "", "", "", "", "", None, "", "", "Weekly", "Wed", D("2026-11-04"), ""],
    ["Deadline", "Fund deposit due", D("2026-11-15"), None, "", "", "", "", "", "", "", "", "", "", None, "", "", "", "", None, ""],
    ["Practice", "Extra serving practice", D("2026-12-12"), None, "10:00 AM", "Test Gym", "", "", "", "", "", "", "", "", None, "", "", "None", "", None, "", "1", "White"],
])
fill(wb["Lists"], [["Packing checklist", "Water bottle"], ["Packing checklist", "Snacks"], ["Uniform items", "Home jersey"],
                    ["Practice uniform colors", "Navy"], ["Practice uniform colors", "White"]])
wb.save(OUT + "team-sample.xlsx")

# ---------- team, with mistakes ----------
wb = load_workbook("web/public/templates/Team-Setup-Template.xlsx")
fill(wb["Roster"], [
    roster("3", "Ivy", "Ruiz", "Kim Ruiz", "not-an-email"),
    ["4", "Max", "Bell", "", "", "", "", "", "", "", ""],
    roster("3", "Zed", "Ray", "Al Ray", "al.ray@example.com"),
])
fill(wb["Schedule"], [
    ["Tournament", "Oops Cup", "next friday", None, "", "", "", "", "", "", "", "", "", "", None, "44", "", "Weekly", "", None, ""],
    ["Party", "Bad type", D("2026-10-01"), None],
])
wb.save(OUT + "team-errors.xlsx")

# ---------- club ----------
wb = load_workbook("web/public/templates/Club-Setup-Template.xlsx")
fill(wb["Club"], [["Test Club", "Test", "test", "#14532D", "#9AD3A5", "Club dues are due in August."]], start=2)
fill(wb["Club links"], [["Club calendar", "https://example.com/calendar"], ["Uniform store", "https://example.com/store"]])
fill(wb["Club admins"], [["Pat", "Lane", "pat.lane@example.com"]])
fill(wb["Teams"], [
    ["test-15a", "Test 15 National", "National", "15U", "15 Open", "2026-27", "", 400, D("2026-11-01"), ""],
    ["test-15b", "Test 15 Regional", "Regional", "15U", "15 USA", "2026-27", "", 300, None, ""],
    ["test-16", "Test 16 National", "National", "16U", "16 Open", "2026-27", "", 400, None, "test-15a"],
])
fill(wb["Staff"], [
    ["test-15a", "Dana", "Fox", "dana.fox@example.com", "555-0201", "Y", "Y", "N", "N", "N", "Y"],
    ["test-15b", "Lee", "Park", "lee.park@example.com", "", "Y", "Y", "", "", "", ""],
    ["test-16", "Dana", "Fox", "dana.fox@example.com", "555-0201", "Y", "Y", "N", "N", "N", "Y"],
])
fill(wb["Rosters"], [
    ["test-15a"] + roster("3", "Ivy", "Ruiz", "Kim Ruiz", "kim.ruiz@example.com"),
    ["test-15a"] + roster("8", "June", "Hale", "Ana Hale", "ana.hale@example.com"),
    ["test-15b"] + roster("5", "Mo", "Diaz", "Lu Diaz", "lu.diaz@example.com"),
    ["test-16"] + roster("11", "Pia", "Ng", "Vi Ng", "vi.ng@example.com"),
])
fill(wb["Practice patterns"], [
    ["15U", "Monday practice", "Monday", "6:00 PM", "8:00 PM", D("2026-10-05"), None, "", "Test Gym", "", "", ""],
    ["test-16", "Thursday practice", "Thursday", "7:00 PM", "9:00 PM", D("2026-10-08"), None, "Every 3 weeks", "Test Gym", "2", "Red", ""],
])
fill(wb["Shared schedule"], [
    ["All National", "Tournament", "National Qualifier", D("2027-02-13"), D("2027-02-15"), "", "Test Expo", "Orlando, FL", "", "Y",
     "", "", "Test Hotel", "https://example.com/hotel", "QUAL27", D("2027-01-10"), "None", "", None, ""],
    ["All", "Deadline", "Club fees due", D("2026-08-31"), None, "", "", "", "", "", "", "", "", "", "", None, "", "", None, ""],
])
fill(wb["Defaults"], [["Handbook section", "Playing time", "Playing time is earned at practice."], ["Packing checklist", "", "Knee pads"],
                       ["Practice uniform colors", "", "Black"]])
wb.save(OUT + "club-sample.xlsx")
print("ok")
